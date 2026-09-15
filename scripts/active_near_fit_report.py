"""Write an `active_near_fit` sheet: still-open postings worth applying to now.

Joins manual_review (decision, should_be_filtered) with jobs_master
(availability_status, url, location) — the same join shape as
skills_gap_report.py — and keeps only postings that are both near-fit
(Подходит/Возможно, not filtered) and still live (availability_status ==
"active"). Point-in-time: availability_status only reflects the last
--check-availability run, not the current moment.
"""

from __future__ import annotations

import sys
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

import openpyxl
from openpyxl.styles import Font, PatternFill

SHEET_NAME = "active_near_fit"
COLUMNS = (
    "decision", "title", "company", "source", "city_region", "country",
    "work_mode", "seniority_manual", "main_reason", "review_notes", "url",
)
DECISION_FILLS = {
    "Подходит": PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
    "Возможно": PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"),
}


def _sheet_rows(worksheet) -> list[dict[str, str]]:
    headers = [cell.value for cell in worksheet[1]]
    rows = []
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        rows.append({header: value for header, value in zip(headers, row) if header})
    return rows


def build_active_near_fit(joined: list[dict[str, str]]) -> list[dict[str, str]]:
    kept = [
        row for row in joined
        if str(row.get("decision", "")).strip() in {"Подходит", "Возможно"}
        and str(row.get("should_be_filtered", "")).strip() != "Да"
        and str(row.get("availability_status", "")).strip() == "active"
    ]
    kept.sort(key=lambda row: (row.get("decision") != "Подходит", row.get("source", "")))
    return kept


def _write_sheet(workbook, rows: list[dict[str, str]]) -> None:
    if SHEET_NAME in workbook.sheetnames:
        del workbook[SHEET_NAME]
    sheet = workbook.create_sheet(SHEET_NAME)
    sheet.append(list(COLUMNS))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append([row.get(column, "") for column in COLUMNS])
        fill = DECISION_FILLS.get(row.get("decision"))
        if fill:
            for cell in sheet[sheet.max_row]:
                cell.fill = fill
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column_cells in sheet.columns:
        width = min(max((len(str(cell.value or "")) for cell in column_cells), default=10) + 2, 50)
        sheet.column_dimensions[column_cells[0].column_letter].width = width


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    path = PROJECT_ROOT / "data" / "processed" / "crypto_jobs_clean_v1.xlsx"
    workbook = openpyxl.load_workbook(path)
    master_by_id = {row.get("job_id"): row for row in _sheet_rows(workbook["jobs_master"])}
    review_rows = _sheet_rows(workbook["manual_review"])

    joined = []
    for review in review_rows:
        master = master_by_id.get(review.get("job_id"), {})
        joined.append({**master, **review})

    rows = build_active_near_fit(joined)
    print(f"Active near-fit postings: {len(rows)}")
    for row in rows:
        print(f"  [{row['decision']}] {row['title']} — {row['company']} ({row['source']})")

    _write_sheet(workbook, rows)
    workbook.save(path)
    print(f"\nSaved '{SHEET_NAME}' sheet in {path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
