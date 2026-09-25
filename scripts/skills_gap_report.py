"""Print a skill-gap and rejection-reason breakdown from the reviewed dataset."""

from __future__ import annotations

import sys
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill

from core.skill_recommendations import (
    build_recommendations,
    count_certification_mentions,
    count_market_demand,
    dedupe_rows,
    format_recommendation,
    posting_text,
)
from core.skills_gap import build_rejection_report, build_skill_gap_report

# A skill mentioned in only one near-fit posting could just be that one
# employer's quirk; recurring across at least this many postings is a real,
# prioritizable signal worth actually studying.
WORTH_LEARNING_MIN_DEMAND = 2
WORTH_LEARNING_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")


def _sheet_rows(worksheet) -> list[dict[str, str]]:
    headers = [cell.value for cell in worksheet[1]]
    rows = []
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        rows.append({header: value for header, value in zip(headers, row) if header})
    return rows


def _bar_chart(sheet, *, title: str, header_row: int, first_row: int, last_row: int) -> BarChart | None:
    """Horizontal bar chart over a (label, count) table already written to the sheet."""
    if last_row < first_row:
        return None
    chart = BarChart()
    chart.type = "bar"  # horizontal — skill/bucket names stay readable
    chart.title = title
    # For a horizontal bar chart, openpyxl's x_axis is the category axis
    # (skill/bucket names) and y_axis is the value axis (count) despite the
    # visual layout being rotated - a "count" title on x_axis (as this used
    # to be) mislabels the category axis instead of the value axis.
    chart.x_axis.title = None
    chart.x_axis.tickLblPos = "nextTo"  # always show category names, never rely on Excel's default
    chart.y_axis.title = "count"
    chart.legend = None
    data = Reference(sheet, min_col=2, min_row=header_row, max_row=last_row)
    categories = Reference(sheet, min_col=1, min_row=first_row, max_row=last_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    chart.height = 7.5
    chart.width = 15
    return chart


def _write_report_sheet(workbook, skill_report, rejection_report) -> None:
    if "skills_gap" in workbook.sheetnames:
        del workbook["skills_gap"]
    sheet = workbook.create_sheet("skills_gap")
    sheet.append(["Skills in near-fit postings I don't have", "count", "example postings"])
    gap_header_row = sheet.max_row
    for skill, count in skill_report.gap.most_common():
        sheet.append([skill, count, "; ".join(skill_report.examples.get(skill, []))])
        if count >= WORTH_LEARNING_MIN_DEMAND:
            for cell in sheet[sheet.max_row]:
                cell.fill = WORTH_LEARNING_FILL
    gap_last_row = sheet.max_row
    sheet.append([])
    sheet.append(["Skills I have at a basic level but postings want more depth", "count", "total demand", "example postings"])
    for skill, count in skill_report.partial.most_common():
        sheet.append([skill, count, skill_report.demand[skill], "; ".join(skill_report.examples.get(skill, []))])
    sheet.append([])
    sheet.append(["Rejection reason (learnable)", "count", "share of rejected %", "example reasons"])
    learnable_header_row = sheet.max_row
    total = rejection_report.total_rejected or 1
    for bucket, count in rejection_report.learnable_counts().most_common():
        sheet.append([bucket, count, round(count / total * 100, 1), " | ".join(rejection_report.examples.get(bucket, []))])
    learnable_last_row = sheet.max_row
    sheet.append([])
    sheet.append(["Rejection reason (structural)", "count", "share of rejected %", "example reasons"])
    learnable_keys = set(rejection_report.learnable_counts())
    structural = {b: c for b, c in rejection_report.counts.items() if b not in learnable_keys}
    for bucket, count in sorted(structural.items(), key=lambda item: -item[1]):
        sheet.append([bucket, count, round(count / total * 100, 1), " | ".join(rejection_report.examples.get(bucket, []))])
    for column_cells in sheet.columns:
        width = min(max((len(str(cell.value or "")) for cell in column_cells), default=10) + 2, 50)
        sheet.column_dimensions[column_cells[0].column_letter].width = width

    # Charts make "where should I focus" visible at a glance instead of
    # requiring a read through every row — placed well past the data columns
    # (A-D) so they never overlap the tables above.
    skill_chart = _bar_chart(
        sheet,
        title="Skills to prioritize (demand across near-fit postings)",
        header_row=gap_header_row,
        first_row=gap_header_row + 1,
        last_row=gap_last_row,
    )
    if skill_chart:
        sheet.add_chart(skill_chart, "G1")

    rejection_chart = _bar_chart(
        sheet,
        title="Why learnable rejections happen (count)",
        header_row=learnable_header_row,
        first_row=learnable_header_row + 1,
        last_row=learnable_last_row,
    )
    if rejection_chart:
        sheet.add_chart(rejection_chart, "G20")


def _write_recommendations_sheet(workbook, recommendations, market_total: int) -> None:
    if "skill_recommendations" in workbook.sheetnames:
        del workbook["skill_recommendations"]
    sheet = workbook.create_sheet("skill_recommendations")
    sheet.append([
        "skill", "status", "near-fit postings", f"market postings (of {market_total})", "market share %",
        "certificate postings", "action", "certificate advice", "note",
    ])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for rec in recommendations:
        sheet.append([
            rec.skill, rec.status, rec.near_fit_postings, rec.market_postings, round(rec.market_share * 100, 1),
            rec.cert_postings, rec.action, rec.certificate_advice, rec.note,
        ])
        if rec.action == "learn":
            for cell in sheet[sheet.max_row]:
                cell.fill = WORTH_LEARNING_FILL
    for column_cells in sheet.columns:
        width = min(max((len(str(cell.value or "")) for cell in column_cells), default=10) + 2, 60)
        sheet.column_dimensions[column_cells[0].column_letter].width = width
    sheet.freeze_panes = "A2"


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    path = PROJECT_ROOT / "data" / "processed" / "crypto_jobs_clean_v1.xlsx"
    workbook = openpyxl.load_workbook(path)
    master_by_id = {row.get("job_id"): row for row in _sheet_rows(workbook["jobs_master"])}
    review_rows = _sheet_rows(workbook["manual_review"])

    joined = []
    for review in review_rows:
        master = master_by_id.get(review.get("job_id"), {})
        joined.append({**master, **review})

    skill_report = build_skill_gap_report(joined)
    rejection_report = build_rejection_report(joined)

    print(f"Reviewed postings: {len(review_rows)}")
    print(f"Rejected/filtered: {rejection_report.total_rejected}\n")

    print("=== Skills that show up in near-fit postings (Подходит/Возможно) but I don't have ===")
    if not skill_report.gap:
        print("(none — every required/preferred skill in near-fit postings is already covered)")
    for skill, count in skill_report.gap.most_common(15):
        examples = ", ".join(skill_report.examples.get(skill, []))
        print(f"  {count:>3}  {skill}  [{examples}]")

    if skill_report.partial:
        print("\n=== Skills I have at a basic level, but postings keep asking for more depth ===")
        for skill, count in skill_report.partial.most_common(10):
            examples = ", ".join(skill_report.examples.get(skill, []))
            print(f"  {count:>3}  {skill}  (demand: {skill_report.demand[skill]})  [{examples}]")

    print("\n=== Why postings get rejected (learnable — worth investing time) ===")
    learnable = rejection_report.learnable_counts()
    if not learnable:
        print("(no learnable-gap rejections found)")
    for bucket, count in learnable.most_common():
        share = count / rejection_report.total_rejected * 100 if rejection_report.total_rejected else 0
        print(f"  {count:>3} ({share:4.1f}%)  {bucket}")
        for example in rejection_report.examples.get(bucket, []):
            print(f"           - {example}")

    print("\n=== Why postings get rejected (structural — not fixable by studying) ===")
    structural = {bucket: count for bucket, count in rejection_report.counts.items() if bucket not in learnable}
    for bucket, count in sorted(structural.items(), key=lambda item: -item[1]):
        share = count / rejection_report.total_rejected * 100 if rejection_report.total_rejected else 0
        print(f"  {count:>3} ({share:4.1f}%)  {bucket}")
        for example in rejection_report.examples.get(bucket, []):
            print(f"           - {example}")

    # Market side: every collected posting (not only near-fit ones), one per
    # normalized company + title, so a skill's share reflects the whole market
    # sample rather than the postings already hand-picked as plausible.
    market_rows = dedupe_rows(master_by_id.values())
    market_texts = [posting_text(row) for row in market_rows]
    studied = set(skill_report.gap) | set(skill_report.partial)
    recommendations = build_recommendations(
        skill_report,
        count_market_demand(market_texts, studied),
        count_certification_mentions(market_texts, studied),
        len(market_rows),
    )
    print(f"\n=== What to study (market demand across {len(market_rows)} unique postings) ===")
    for rec in recommendations:
        print("  " + format_recommendation(rec))

    _write_report_sheet(workbook, skill_report, rejection_report)
    _write_recommendations_sheet(workbook, recommendations, len(market_rows))
    workbook.save(path)
    print(f"\nSaved 'skills_gap' and 'skill_recommendations' sheets in {path}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
