"""Write/refresh the `applications` tracking sheet from `CV/drafts/*.md`
(plan item 8.5).

No LLM call, no API key, no network access - purely local bookkeeping.
Every drafted posting starts as `status=draft`; a human moves a row to
`approved` (or `rejected`) by hand, and only to `sent` after actually
sending it themselves outside this tool. Rerunning this script never
resets an existing `status`/`notes` back to defaults - see
`core/applications.py:build_applications_rows`.
"""

from __future__ import annotations

import sys
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

import pandas as pd
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from core.applications import (
    APPLICATIONS_COLUMNS,
    APPLICATIONS_INPUT_COLUMNS,
    STATUS_OPTIONS,
    build_applications_rows,
    scan_draft_files,
)
from core.storage import load_records

DATA_PATH = PROJECT_ROOT / "data" / "processed" / "crypto_jobs_clean_v1.xlsx"
DEFAULT_DRAFTS_DIR = PROJECT_ROOT / "CV" / "drafts"
SHEET_NAME = "applications"

STATUS_FILLS = {
    "draft": PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid"),
    "approved": PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
    "rejected": PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
    "sent": PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid"),
}


def load_existing_applications(path: Path) -> dict[str, dict[str, str]]:
    if not path.exists():
        return {}
    try:
        with pd.ExcelFile(path) as workbook:
            if SHEET_NAME not in workbook.sheet_names:
                return {}
            dataframe = pd.read_excel(workbook, sheet_name=SHEET_NAME)
    except (OSError, ValueError):
        return {}
    if "job_id" not in dataframe.columns:
        return {}
    result = {}
    for _, row in dataframe.iterrows():
        job_id = str(row.get("job_id", "")).strip()
        if job_id and job_id.casefold() != "nan":
            result[job_id] = {
                column: ("" if isinstance(value, float) and value != value else value)
                for column, value in row.to_dict().items()
            }
    return result


def write_applications_sheet(path: Path, rows: list[dict[str, str]]) -> None:
    dataframe = pd.DataFrame(rows, columns=APPLICATIONS_COLUMNS)
    writer_kwargs = {"engine": "openpyxl"}
    if path.exists():
        writer_kwargs["mode"] = "a"
        writer_kwargs["if_sheet_exists"] = "replace"
    with pd.ExcelWriter(path, **writer_kwargs) as writer:
        dataframe.to_excel(writer, sheet_name=SHEET_NAME, index=False)
        sheet = writer.book[SHEET_NAME]
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column_cells in sheet.columns:
            values = [str(cell.value or "") for cell in column_cells[:100]]
            width = min(max(max((len(value) for value in values), default=10) + 2, 10), 50)
            sheet.column_dimensions[column_cells[0].column_letter].width = width

        header_yellow = PatternFill("solid", fgColor="FFD966")
        yellow = PatternFill("solid", fgColor="FFF2CC")
        for column_name in APPLICATIONS_INPUT_COLUMNS:
            column_index = APPLICATIONS_COLUMNS.index(column_name) + 1
            sheet.cell(1, column_index).fill = header_yellow
            sheet.cell(1, column_index).font = Font(bold=True)
            for row_index in range(2, sheet.max_row + 1):
                sheet.cell(row_index, column_index).fill = yellow

        status_column_letter = sheet.cell(1, APPLICATIONS_COLUMNS.index("status") + 1).column_letter
        validation = DataValidation(type="list", formula1=f'"{",".join(STATUS_OPTIONS)}"', allow_blank=False)
        sheet.add_data_validation(validation)
        validation.add(f"{status_column_letter}2:{status_column_letter}{max(sheet.max_row, 2)}")

        last_column_letter = sheet.cell(1, len(APPLICATIONS_COLUMNS)).column_letter
        data_range = f"A2:{last_column_letter}{max(sheet.max_row, 2)}"
        for status, fill in STATUS_FILLS.items():
            sheet.conditional_formatting.add(
                data_range,
                FormulaRule(formula=[f'${status_column_letter}2="{status}"'], fill=fill),
            )


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")

    records_by_id = {record.job_id: record for record in load_records(DATA_PATH)}
    drafts = scan_draft_files(DEFAULT_DRAFTS_DIR)
    existing = load_existing_applications(DATA_PATH)
    rows = build_applications_rows(drafts, records_by_id, existing)

    if not rows:
        print(f"No draft files found in {DEFAULT_DRAFTS_DIR}")
        return 0

    write_applications_sheet(DATA_PATH, rows)

    print(f"Applications tracked: {len(rows)}")
    for row in rows:
        print(f"  [{row['status']:>8}] {row['title']} — {row['company']} (review: {row['review_verdict'] or 'none'}, pdf: {row['pdf_ready']})")
    print(f"Saved '{SHEET_NAME}' sheet in {DATA_PATH}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
