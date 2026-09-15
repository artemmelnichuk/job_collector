from __future__ import annotations

import json
import shutil
import unittest
from pathlib import Path

from openpyxl import load_workbook

from core.logging import RunStats, write_run_log
from core.models import JobRecord
from core.storage import compute_personal_fit, load_manual_review, load_records, merge_records, save_records
from scripts.collector import ats_job_is_present, detect_generic_availability, parse_ats_board


def make_record(**overrides: str) -> JobRecord:
    values = {
        "source": "LinkedIn",
        "title": "Trading Analyst",
        "company": "Example Exchange",
        "city_region": "Paris",
        "url": "https://example.com/jobs/123",
        "full_text": "SQL and market data",
    }
    values.update(overrides)
    return JobRecord(**values)


class StorageTests(unittest.TestCase):
    def test_save_and_load_round_trip(self) -> None:
        path = Path("data/raw/_test_storage_run.xlsx")
        try:
            save_records([make_record()], path)

            loaded = load_records(path)

            self.assertEqual(len(loaded), 1)
            self.assertEqual(loaded[0].title, "Trading Analyst")
            self.assertTrue(path.exists())
            self.assertTrue(path.with_suffix(".csv").exists())
            workbook = load_workbook(path, read_only=True)
            self.assertEqual(workbook.sheetnames, ["jobs_master", "jobs_view", "jobs_text", "manual_review"])
            workbook.close()
        finally:
            path.unlink(missing_ok=True)
            path.with_suffix(".csv").unlink(missing_ok=True)

    def test_manual_review_shows_category_and_preserves_should_be_filtered(self) -> None:
        path = Path("data/raw/_test_manual_review_run.xlsx")
        try:
            record = make_record(job_id="linkedin_test123", job_category="crypto")
            save_records([record], path)

            workbook = load_workbook(path)
            sheet = workbook["manual_review"]
            headers = [cell.value for cell in sheet[1]]
            self.assertIn("job_category", headers)
            self.assertIn("should_be_filtered", headers)
            row = dict(zip(headers, (cell.value for cell in sheet[2])))
            self.assertEqual(row["job_category"], "crypto")

            flag_column = headers.index("should_be_filtered") + 1
            sheet.cell(2, flag_column).value = "Да"
            workbook.save(path)
            workbook.close()

            save_records([record], path)

            reloaded = load_workbook(path)
            reloaded_sheet = reloaded["manual_review"]
            reloaded_headers = [cell.value for cell in reloaded_sheet[1]]
            reloaded_row = dict(zip(reloaded_headers, (cell.value for cell in reloaded_sheet[2])))
            self.assertEqual(reloaded_row["should_be_filtered"], "Да")
            self.assertEqual(reloaded_row["job_category"], "crypto")
            reloaded.close()
        finally:
            path.unlink(missing_ok=True)
            path.with_suffix(".csv").unlink(missing_ok=True)

    def test_save_records_preserves_sheets_it_does_not_own(self) -> None:
        # Regression: save_records used to rebuild the workbook from scratch
        # on every call, silently deleting any sheet it doesn't itself write
        # (e.g. skills_gap_report.py's "skills_gap" sheet) on every later
        # analyzer.py run or --check-availability pass.
        path = Path("data/raw/_test_preserves_extra_sheet.xlsx")
        try:
            save_records([make_record(job_id="preserve_test")], path)

            workbook = load_workbook(path)
            workbook.create_sheet("skills_gap").append(["some", "report", "data"])
            workbook.save(path)
            workbook.close()

            save_records([make_record(job_id="preserve_test")], path)

            reloaded = load_workbook(path, read_only=True)
            self.assertIn("skills_gap", reloaded.sheetnames)
            reloaded.close()
        finally:
            path.unlink(missing_ok=True)
            path.with_suffix(".csv").unlink(missing_ok=True)

    def test_compute_personal_fit_folds_decision_and_reason(self) -> None:
        self.assertEqual(compute_personal_fit("Подходит", "irrelevant"), "Подходит")
        self.assertEqual(compute_personal_fit("Возможно", "нужно больше опыта"), "Возможно — нужно больше опыта")
        self.assertEqual(compute_personal_fit("Не подходит", "офис в Украине"), "Не подходит — офис в Украине")
        self.assertEqual(compute_personal_fit("Возможно", ""), "Возможно")
        self.assertEqual(compute_personal_fit("", ""), "")

    def test_load_manual_review_returns_empty_string_not_nan_for_blank_cells(self) -> None:
        # Regression: pd.read_excel reads an untouched blank cell as float
        # NaN, not "". Left unconverted, str(nan) is the literal text "nan" -
        # truthy and non-empty - which broke both compute_personal_fit
        # (folded into "nan — nan" for postings nobody reviewed yet) and
        # scripts/fit_score.py's undecided-postings filter (treated "nan" as
        # an already-made decision and skipped scoring brand-new postings).
        path = Path("data/raw/_test_manual_review_nan_cells.xlsx")
        try:
            record = make_record(job_id="linkedin_never_reviewed")
            save_records([record], path)
            # A second save (e.g. a later analyzer.py rerun) is what actually
            # exercises the bug: it reads the still-blank cells back off disk
            # via load_manual_review/pd.read_excel (NaN), not the in-memory
            # "" the first save started from.
            save_records([record], path)

            manual = load_manual_review(path)
            row = manual["linkedin_never_reviewed"]
            self.assertEqual(row["decision"], "")
            self.assertEqual(row["main_reason"], "")

            workbook = load_workbook(path)
            headers = [cell.value for cell in workbook["manual_review"][1]]
            personal_fit = dict(zip(headers, (cell.value for cell in workbook["manual_review"][2])))["personal_fit"]
            # openpyxl reads a written-as-"" cell back as None, not "" - that's
            # normal Excel round-tripping, not the bug under test here. The
            # bug this test guards against is the literal text "nan — nan".
            self.assertIn(personal_fit, (None, ""))
            workbook.close()
        finally:
            path.unlink(missing_ok=True)
            path.with_suffix(".csv").unlink(missing_ok=True)

    def test_manual_review_personal_fit_is_recomputed_from_saved_decision(self) -> None:
        path = Path("data/raw/_test_manual_review_personal_fit.xlsx")
        try:
            record = make_record(job_id="linkedin_personal_fit_test")
            save_records([record], path)

            workbook = load_workbook(path)
            sheet = workbook["manual_review"]
            headers = [cell.value for cell in sheet[1]]
            decision_col = headers.index("decision") + 1
            reason_col = headers.index("main_reason") + 1
            sheet.cell(2, decision_col).value = "Не подходит"
            sheet.cell(2, reason_col).value = "language_other_than_english"
            workbook.save(path)
            workbook.close()

            # Recompute on the next save (e.g. after an analyzer rerun) - not
            # just left over from the manual edit above.
            save_records([record], path)

            reloaded = load_workbook(path)
            reloaded_sheet = reloaded["manual_review"]
            reloaded_headers = [cell.value for cell in reloaded_sheet[1]]
            reloaded_row = dict(zip(reloaded_headers, (cell.value for cell in reloaded_sheet[2])))
            self.assertEqual(reloaded_row["personal_fit"], "Не подходит — language_other_than_english")
            reloaded.close()
        finally:
            path.unlink(missing_ok=True)
            path.with_suffix(".csv").unlink(missing_ok=True)

    def test_manual_review_has_conditional_formatting_for_each_decision(self) -> None:
        path = Path("data/raw/_test_manual_review_colors.xlsx")
        try:
            save_records([make_record(job_id="linkedin_color_test")], path)

            workbook = load_workbook(path)
            sheet = workbook["manual_review"]
            formulas = {
                formula
                for cf_range in sheet.conditional_formatting
                for rule in cf_range.rules
                for formula in rule.formula
            }
            self.assertIn('$M2="Подходит"', formulas)
            self.assertIn('$M2="Возможно"', formulas)
            self.assertIn('$M2="Не подходит"', formulas)
            workbook.close()
        finally:
            path.unlink(missing_ok=True)
            path.with_suffix(".csv").unlink(missing_ok=True)

    def test_merge_classifies_new_existing_and_updated(self) -> None:
        original = make_record()
        merged, first_counts = merge_records([], [original])
        self.assertEqual(first_counts, {"new": 1, "existing": 0, "updated": 0})

        same, same_counts = merge_records(merged, [make_record()])
        self.assertEqual(same_counts, {"new": 0, "existing": 1, "updated": 0})
        self.assertEqual(same[0].job_id, merged[0].job_id)

        _, updated_counts = merge_records(merged, [make_record(full_text="Updated text")])
        self.assertEqual(updated_counts, {"new": 0, "existing": 0, "updated": 1})

    def test_recollecting_a_posting_does_not_erase_prior_analysis(self) -> None:
        analyzed = make_record(
            role_family="Data Analytics",
            seniority="Junior",
            skills_all="SQL; Python",
            required_skills="SQL",
            preferred_skills="Python",
            analysis_note="Rule-based first pass; review before using for decisions.",
            availability_status="active",
        )
        merged, _ = merge_records([], [analyzed])

        # A plain re-collect never carries analysis fields (analyzer.py runs
        # separately) and always sets availability_status="active" by default.
        rescraped = make_record(availability_status="active")
        remerged, counts = merge_records(merged, [rescraped])

        self.assertEqual(counts, {"new": 0, "existing": 1, "updated": 0})
        self.assertEqual(remerged[0].role_family, "Data Analytics")
        self.assertEqual(remerged[0].skills_all, "SQL; Python")
        self.assertEqual(remerged[0].analysis_note, "Rule-based first pass; review before using for decisions.")

    def test_recollecting_does_not_downgrade_a_checked_availability_status(self) -> None:
        checked_closed = make_record(availability_status="closed")
        merged, _ = merge_records([], [checked_closed])

        rescraped = make_record(availability_status="active")
        remerged, _ = merge_records(merged, [rescraped])

        self.assertEqual(remerged[0].availability_status, "closed")

    def test_a_real_content_change_still_reports_as_updated_and_keeps_analysis(self) -> None:
        analyzed = make_record(role_family="Data Analytics", availability_status="active")
        merged, _ = merge_records([], [analyzed])

        changed = make_record(full_text="Updated text", availability_status="active")
        remerged, counts = merge_records(merged, [changed])

        self.assertEqual(counts, {"new": 0, "existing": 0, "updated": 1})
        self.assertEqual(remerged[0].full_text, "Updated text")
        self.assertEqual(remerged[0].role_family, "Data Analytics")


class LoggingTests(unittest.TestCase):
    def test_report_contains_source_and_totals(self) -> None:
        stats = RunStats()
        source = stats.for_source("LinkedIn")
        source.queries = ["Trading Analyst"]
        source.found = 3
        source.unique = 2
        source.duplicates = 1
        source.filtered_out = 1
        source.errors = 0

        report = "\n".join(stats.report_lines())

        self.assertIn("Source: LinkedIn", report)
        self.assertIn("Total found: 3", report)
        self.assertIn("Unique jobs: 2", report)
        self.assertIn("Filtered out (off-context or overqualified): 1", report)

    def test_to_dict_serializes_every_source(self) -> None:
        stats = RunStats()
        stats.for_source("LinkedIn").found = 3
        stats.for_source("Djinni").found = 1
        as_dict = stats.to_dict()
        self.assertEqual(as_dict["LinkedIn"]["found"], 3)
        self.assertEqual(as_dict["Djinni"]["found"], 1)


class RunLogTests(unittest.TestCase):
    def test_write_run_log_appends_a_json_line_and_creates_the_directory(self) -> None:
        logs_dir = Path("data/raw/_test_logs")
        shutil.rmtree(logs_dir, ignore_errors=True)
        try:
            log_path = write_run_log(logs_dir, "collect", {"sources": {"LinkedIn": {"found": 3}}})
            self.assertTrue(log_path.exists())
            lines = log_path.read_text(encoding="utf-8").strip().splitlines()
            self.assertEqual(len(lines), 1)
            record = json.loads(lines[0])
            self.assertEqual(record["run_type"], "collect")
            self.assertEqual(record["sources"]["LinkedIn"]["found"], 3)
            self.assertIn("timestamp", record)

            write_run_log(logs_dir, "collect", {"sources": {}})
            lines = log_path.read_text(encoding="utf-8").strip().splitlines()
            self.assertEqual(len(lines), 2)
        finally:
            shutil.rmtree(logs_dir, ignore_errors=True)


class AvailabilityTests(unittest.TestCase):
    def test_generic_availability_detects_http_closed_page(self) -> None:
        self.assertEqual(detect_generic_availability(404, "Not found"), "closed")
        self.assertEqual(detect_generic_availability(200, "Open role"), "active")
        self.assertEqual(detect_generic_availability(None, ""), "unknown")

    def test_ats_reference_and_job_matching(self) -> None:
        self.assertEqual(parse_ats_board("ATS: Greenhouse; board: coinbase"), ("greenhouse", "coinbase"))
        jobs = [{"id": 123, "absolute_url": "https://www.coinbase.com/careers/positions/123?gh_jid=123"}]
        self.assertTrue(ats_job_is_present("https://www.coinbase.com/careers/positions/123?gh_jid=123", "greenhouse", jobs))
        self.assertFalse(ats_job_is_present("https://www.coinbase.com/careers/positions/999?gh_jid=999", "greenhouse", jobs))
