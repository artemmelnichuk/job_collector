from __future__ import annotations

import unittest
from pathlib import Path

from openpyxl import load_workbook

from scripts.applications_tracker import load_existing_applications, write_applications_sheet


class ApplicationsSheetRoundTripTests(unittest.TestCase):
    def test_write_then_load_round_trips_status_and_notes(self) -> None:
        path = Path("data/raw/_test_applications_sheet.xlsx")
        try:
            rows = [
                {
                    "job_id": "job_a",
                    "title": "Data Analyst",
                    "company": "Acme",
                    "url": "https://example.com/a",
                    "review_verdict": "approved",
                    "pdf_ready": "Да",
                    "status": "draft",
                    "notes": "",
                    "date_added": "2026-09-15",
                }
            ]
            write_applications_sheet(path, rows)

            workbook = load_workbook(path)
            self.assertIn("applications", workbook.sheetnames)
            sheet = workbook["applications"]
            headers = [cell.value for cell in sheet[1]]
            self.assertIn("status", headers)
            status_col = headers.index("status") + 1
            sheet.cell(2, status_col).value = "approved"
            workbook.save(path)
            workbook.close()

            existing = load_existing_applications(path)
            self.assertEqual(existing["job_a"]["status"], "approved")
        finally:
            path.unlink(missing_ok=True)

    def test_load_existing_returns_empty_dict_when_sheet_missing(self) -> None:
        path = Path("data/raw/_test_applications_missing.xlsx")
        try:
            write_applications_sheet(
                path,
                [
                    {
                        "job_id": "x",
                        "title": "",
                        "company": "",
                        "url": "",
                        "review_verdict": "",
                        "pdf_ready": "Нет",
                        "status": "draft",
                        "notes": "",
                        "date_added": "2026-09-15",
                    }
                ],
            )
            # Overwrite with a workbook that has no "applications" sheet at all.
            from openpyxl import Workbook

            wb = Workbook()
            wb.save(path)

            self.assertEqual(load_existing_applications(path), {})
        finally:
            path.unlink(missing_ok=True)

    def test_load_existing_handles_missing_file(self) -> None:
        self.assertEqual(load_existing_applications(Path("data/raw/_does_not_exist.xlsx")), {})


if __name__ == "__main__":
    unittest.main()
